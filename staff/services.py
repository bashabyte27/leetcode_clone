import csv
import io
import json
from datetime import timezone as dt_timezone
from decimal import Decimal

import pandas as pd
from django.contrib.auth.password_validation import validate_password
from django.db import transaction
from django.http import HttpResponse
from django.utils.timezone import localtime
from openpyxl import Workbook

from leaderboard.services import get_user_standing
from problems.models import DifficultyChoices, Problem, TestCase
from submissions.models import Submission, SubmissionStatusChoices, UserSolvedProblem
from users.models import UserProfile, UserStats, Users


def _clean(value):
    if pd.isna(value):
        return ''
    return str(value).strip()


def _read_table(source, sheet_name=None):
    name = getattr(source, 'name', str(source)).lower()
    if name.endswith('.csv') and sheet_name is None:
        return pd.read_csv(source)
    return pd.read_excel(source, sheet_name=sheet_name)


def _ensure_languages():
    from problems.models import Language

    languages = [
        ('Python', 'python', '3.0', '71'), ('Java', 'java', '17.0.6', '62'),
        ('C++', 'cpp', '17.0.6', '54'), ('C', 'c', '10.2.0', '50'),
        ('JavaScript', 'javascript', '18.15.0', '63'), ('Go', 'go', '1.18.0', '39'),
        ('Rust', 'rust', '1.65.0', '73'),
    ]
    for name, slug, version, judge_id in languages:
        Language.objects.get_or_create(slug=slug, defaults={
            'name': name, 'version': version, 'judge_id': judge_id, 'is_active': True,
        })


def import_problems(source, actor=None):
    """Import the same Problems/TestCases workbook contract as the command."""
    _ensure_languages()
    errors, created, updated = [], 0, 0
    try:
        problems_df = _read_table(source, 'Problems')
        testcases_df = None
        if not str(getattr(source, 'name', source)).lower().endswith('.csv'):
            testcases_df = _read_table(source, 'TestCases')
    except Exception as exc:
        return {'created': 0, 'updated': 0, 'errors': [f'Workbook could not be read: {exc}']}

    required = {'order_num', 'title', 'description', 'difficulty'}
    missing = required - set(problems_df.columns)
    if missing:
        return {'created': 0, 'updated': 0, 'errors': [f'Missing Problems columns: {", ".join(sorted(missing))}']}

    for row_number, row in problems_df.iterrows():
        display_row = row_number + 2
        try:
            order_num = int(row['order_num'])
            title = _clean(row['title'])
            difficulty = _clean(row['difficulty']).lower()
            if not title or difficulty not in DifficultyChoices.values:
                raise ValueError('title is required and difficulty must be easy, medium, or hard')
            with transaction.atomic():
                problem, was_created = Problem.objects.update_or_create(
                    order_num=order_num,
                    defaults={'title': title, 'description': _clean(row['description']).replace('\\n', '\n'), 'difficulty': difficulty},
                )
            created += int(was_created)
            updated += int(not was_created)
        except Exception as exc:
            errors.append(f'Problems row {display_row}: {exc}')

    if testcases_df is not None:
        required_tests = {'problem_order_num', 'order_num', 'input_data', 'expected_output', 'is_sample'}
        missing = required_tests - set(testcases_df.columns)
        if missing:
            errors.append(f'Missing TestCases columns: {", ".join(sorted(missing))}')
        else:
            for row_number, row in testcases_df.iterrows():
                display_row = row_number + 2
                try:
                    problem = Problem.objects.get(order_num=int(row['problem_order_num']))
                    TestCase.objects.update_or_create(
                        problem=problem,
                        order_num=int(row['order_num']),
                        defaults={
                            'input_data': _clean(row['input_data']).replace('\\n', '\n'),
                            'expected_output': _clean(row['expected_output']).replace('\\n', '\n'),
                            'is_sample': _clean(row['is_sample']).lower() in {'true', '1', 'yes'},
                            'explanation': _clean(row.get('explanation')) or None,
                        },
                    )
                except Exception as exc:
                    errors.append(f'TestCases row {display_row}: {exc}')
    return {'created': created, 'updated': updated, 'errors': errors}


def import_users(source):
    try:
        frame = _read_table(source)
    except Exception as exc:
        return {'created': 0, 'errors': [f'File could not be read: {exc}']}
    required = {'email', 'user_name', 'password'}
    missing = required - set(frame.columns)
    if missing:
        return {'created': 0, 'errors': [f'Missing columns: {", ".join(sorted(missing))}']}
    created, errors = 0, []
    for row_number, row in frame.iterrows():
        display_row = row_number + 2
        email, user_name, password = _clean(row['email']).lower(), _clean(row['user_name']), _clean(row['password'])
        try:
            if not email or not user_name or not password:
                raise ValueError('email, user_name, and password are required')
            if Users.objects.filter(email__iexact=email).exists():
                raise ValueError('email already exists')
            if Users.objects.filter(user_name__iexact=user_name).exists():
                raise ValueError('user_name already exists')
            user = Users(email=email, user_name=user_name, is_verified=True)
            validate_password(password, user)
            user.set_password(password)
            user.full_clean()
            user.save()
            created += 1
        except Exception as exc:
            errors.append(f'Row {display_row}: {exc}')
    return {'created': created, 'errors': errors}


def user_snapshot(user):
    stats, _ = UserStats.objects.get_or_create(user=user)
    standing = get_user_standing(user.id)
    submissions = Submission.objects.filter(user=user).select_related('problem', 'language').order_by('-submitted_at')
    solved = UserSolvedProblem.objects.filter(user=user).select_related('problem').order_by('-last_solved_at')
    total = submissions.count()
    accepted = submissions.filter(status=SubmissionStatusChoices.ACCEPTED).count()
    return {
        'stats': stats, 'standing': standing, 'submissions': submissions, 'solved': solved,
        'total_submissions': total, 'accepted_submissions': accepted,
        'wrong_submissions': submissions.filter(status=SubmissionStatusChoices.WRONG_ANSWER).count(),
        'acceptance_rate': round(accepted * 100 / total, 2) if total else 0,
    }


def xlsx_response(user, snapshot):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    rows = [
        ('User name', user.user_name), ('Email', user.email), ('Status', 'Active' if user.is_active else 'Blocked'),
        ('Total solved', snapshot['solved'].count()), ('Easy solved', snapshot['standing']['easy_solved']),
        ('Medium solved', snapshot['standing']['medium_solved']), ('Hard solved', snapshot['standing']['hard_solved']),
        ('Total submissions', snapshot['total_submissions']), ('Accepted submissions', snapshot['accepted_submissions']),
        ('Wrong submissions', snapshot['wrong_submissions']), ('Acceptance rate', f"{snapshot['acceptance_rate']}%"),
        ('Points', snapshot['standing']['points']), ('Rank', snapshot['standing']['rank'] or 'Unranked'),
    ]
    for row in rows:
        summary.append(row)
    solved_sheet = workbook.create_sheet('Solved problems')
    solved_sheet.append(['Order', 'Title', 'Difficulty', 'First solved', 'Last solved'])
    for item in snapshot['solved']:
        solved_sheet.append([item.problem.order_num, item.problem.title, item.problem.get_difficulty_display(), item.first_solved_at.isoformat(), item.last_solved_at.isoformat()])
    history = workbook.create_sheet('Submission history')
    history.append(['Submitted at', 'Problem', 'Difficulty', 'Language', 'Status', 'Runtime ms', 'Memory kb'])
    for submission in snapshot['submissions']:
        history.append([submission.submitted_at.isoformat(), submission.problem.title, submission.problem.get_difficulty_display(), submission.language.name if submission.language else '', submission.get_status_display(), submission.runtime_ms, submission.memory_kb])
    output = io.BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{user.user_name}-progress.xlsx"'
    return response


def pdf_response(user, snapshot):
    lines = [
        f'BashaByte progress report: {user.user_name}', f'Email: {user.email}',
        f'Status: {"Active" if user.is_active else "Blocked"}',
        f'Solved: {snapshot["solved"].count()} | Easy {snapshot["standing"]["easy_solved"]} | Medium {snapshot["standing"]["medium_solved"]} | Hard {snapshot["standing"]["hard_solved"]}',
        f'Submissions: {snapshot["total_submissions"]} | Accepted {snapshot["accepted_submissions"]} | Wrong {snapshot["wrong_submissions"]} | Rate {snapshot["acceptance_rate"]}%',
        f'Points: {snapshot["standing"]["points"]} | Rank: {snapshot["standing"]["rank"] or "Unranked"}', '', 'Solved problems:',
    ]
    lines.extend(f'#{item.problem.order_num or "-"} {item.problem.title} [{item.problem.get_difficulty_display()}]' for item in snapshot['solved'][:200])
    lines.extend(['', 'Submission history:'])
    lines.extend(f'{localtime(item.submitted_at):%Y-%m-%d %H:%M} | {item.problem.title} | {item.get_status_display()}' for item in snapshot['submissions'][:200])
    content = ['BT /F1 11 Tf 50 760 Td']
    for index, line in enumerate(lines):
        safe = line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')[:115]
        content.append(f'({safe}) Tj')
        if index < len(lines) - 1:
            content.append('0 -15 Td')
    stream = '\n'.join(content).encode('latin-1', 'replace')
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>', b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>', b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>', b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    pdf = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(pdf)); pdf.extend(f'{number} 0 obj\n'.encode()); pdf.extend(obj); pdf.extend(b'\nendobj\n')
    xref = len(pdf); pdf.extend(f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode())
    pdf.extend(b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:]))
    pdf.extend(f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode())
    response = HttpResponse(bytes(pdf), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{user.user_name}-progress.pdf"'
    return response

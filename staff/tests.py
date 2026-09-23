import io
import json

from openpyxl import load_workbook
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from problems.models import Problem, TestCase as ProblemTestCase
from submissions.models import Submission, SubmissionStatusChoices, UserSolvedProblem
from users.models import Users

from .services import import_users


class StaffAccessTests(TestCase):
    def setUp(self):
        self.user = Users.objects.create_user('user@example.com', 'regular_user', 'Password123!')
        self.staff = Users.objects.create_user('staff@example.com', 'staff_user', 'Password123!', is_staff=True)

    def test_staff_area_requires_staff(self):
        response = self.client.get(reverse('staff:dashboard'))
        self.assertRedirects(response, f'{reverse("users:login")}?next={reverse("staff:dashboard")}')
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('staff:dashboard')).status_code, 302)

    def test_staff_can_view_dashboard_and_problem_list(self):
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse('staff:dashboard')).status_code, 200)
        self.assertEqual(self.client.get(reverse('staff:problem_list')).status_code, 200)

    def test_staff_user_search_matches_usernames_only(self):
        self.client.force_login(self.staff)
        target = Users.objects.create_user('target@example.com', 'target_user', 'Password123!')
        email_match = Users.objects.create_user('different@example.com', 'other_user', 'Password123!')
        response = self.client.get(reverse('staff:user_list'), {'q': 'target'})
        self.assertContains(response, target.user_name)
        self.assertNotContains(response, email_match.user_name)

    def test_staff_problem_search_matches_titles_only(self):
        self.client.force_login(self.staff)
        target = Problem.objects.create(title='Binary Search', slug='binary-search', description='d', difficulty='easy')
        slug_match = Problem.objects.create(title='Other Problem', slug='binary-search-extra', description='d', difficulty='easy')
        response = self.client.get(reverse('staff:problem_list'), {'q': 'Binary Search'})
        self.assertContains(response, target.title)
        self.assertNotContains(response, slug_match.title)

    def test_block_action_requires_post(self):
        self.client.force_login(self.staff)
        target = Users.objects.create_user('other@example.com', 'other_user', 'Password123!')
        response = self.client.get(reverse('staff:user_toggle_active', args=[target.id]))
        target.refresh_from_db()
        self.assertEqual(response.status_code, 302)
        self.assertTrue(target.is_active)

    def test_user_import_reports_duplicate_rows(self):
        upload = SimpleUploadedFile(
            'users.csv', b'email,user_name,password\nnew@example.com,new_user,StrongPassword123!\nuser@example.com,duplicate,StrongPassword123!\n',
            content_type='text/csv',
        )
        result = import_users(upload)
        self.assertEqual(result['created'], 1)
        self.assertEqual(len(result['errors']), 1)

    def test_progress_export_is_xlsx(self):
        self.client.force_login(self.staff)
        target = Users.objects.create_user('export@example.com', 'export_user', 'Password123!')
        problem = Problem.objects.create(title='Export problem', description='A problem', difficulty='easy')
        Submission.objects.create(user=target, problem=problem, code='print(1)', status=SubmissionStatusChoices.ACCEPTED)
        response = self.client.get(reverse('staff:user_export_xlsx', args=[target.id]))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Summary', 'Solved problems', 'Submission history'])

    def test_problem_create_generates_slug_and_testcases(self):
        self.client.force_login(self.staff)
        response = self.client.post(reverse('staff:problem_create'), {
            'title': 'Two Sum Practice',
            'description': 'Return the matching pair.',
            'difficulty': 'easy',
            'test_cases': '[{"input_data":"[2,7]","expected_output":"9","is_sample":true}]',
        })
        problem = Problem.objects.get(title='Two Sum Practice')
        self.assertRedirects(response, reverse('staff:problem_list'))
        self.assertEqual(problem.slug, 'two-sum-practice')
        self.assertEqual(ProblemTestCase.objects.filter(problem=problem).count(), 1)

    def test_problem_create_preserves_multiline_expected_output(self):
        self.client.force_login(self.staff)
        expected = '11 x 1 = 11\n11 x 2 = 22\n11 x 3 = 33\n11 x 4 = 44'
        self.client.post(reverse('staff:problem_create'), {
            'title': 'Multiline output problem',
            'description': 'Print each line.',
            'difficulty': 'easy',
            'test_cases': json.dumps([{
                'input_data': '11',
                'expected_output': expected,
                'is_sample': True,
            }]),
        })
        case = ProblemTestCase.objects.get(problem__title='Multiline output problem')
        self.assertEqual(case.expected_output, expected)

    def test_invalid_testcases_do_not_create_problem(self):
        self.client.force_login(self.staff)
        response = self.client.post(reverse('staff:problem_create'), {
            'title': 'Broken Testcase Problem',
            'description': 'This should fail.',
            'difficulty': 'medium',
            'test_cases': '[{"input_data":"only input"}]',
        })
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Problem.objects.filter(title='Broken Testcase Problem').exists())

    def test_user_delete_is_post_only_and_requires_request(self):
        self.client.force_login(self.staff)
        target = Users.objects.create_user('delete@example.com', 'delete_user', 'Password123!')
        self.client.get(reverse('staff:user_delete', args=[target.id]))
        self.assertTrue(Users.objects.filter(id=target.id).exists())
        self.client.post(reverse('staff:user_delete', args=[target.id]))
        self.assertFalse(Users.objects.filter(id=target.id).exists())

    def test_staff_can_view_user_progress_and_submission_detail(self):
        self.client.force_login(self.staff)
        target = Users.objects.create_user('inspect@example.com', 'inspect_user', 'Password123!')
        problem = Problem.objects.create(title='Inspect problem', description='d', difficulty='easy')
        submission = Submission.objects.create(user=target, problem=problem, code='x', status=SubmissionStatusChoices.ACCEPTED)
        UserSolvedProblem.objects.create(user=target, problem=problem, best_submission=submission, first_solved_at=submission.submitted_at, last_solved_at=submission.submitted_at)
        page = self.client.get(reverse('staff:user_detail', args=[target.id]))
        self.assertContains(page, 'Inspect problem')
        detail = self.client.get(
            reverse('submissions:submission_detail', args=[submission.id]),
            HTTP_X_STAFF_INSPECTION='1',
        )
        self.assertEqual(detail.json()['problem'], 'Inspect problem')

    def test_staff_profile_shows_edit_actions_but_hides_account_form_until_edit_is_clicked(self):
        self.client.force_login(self.staff)
        target = Users.objects.create_user('profile@example.com', 'profile_user', 'Password123!')

        default_page = self.client.get(reverse('staff:user_detail', args=[target.id]))
        self.assertContains(default_page, 'Edit User')
        self.assertContains(default_page, 'Delete User')
        self.assertNotContains(default_page, 'Save account changes')

        edit_page = self.client.get(reverse('staff:user_detail', args=[target.id]), {'edit': '1'})
        self.assertContains(edit_page, 'Save account changes')
        self.assertContains(edit_page, 'Staff account management')

    def test_staff_can_edit_manage_test_cases_and_delete_problem(self):
        self.client.force_login(self.staff)
        problem = Problem.objects.create(title='Manage problem', description='old', difficulty='easy')
        case = ProblemTestCase.objects.create(problem=problem, order_num=1, input_data='1', expected_output='1')
        response = self.client.post(reverse('staff:problem_detail', args=[problem.id]), {
            'title': 'Managed problem', 'description': 'new', 'difficulty': 'medium',
            'test_cases': json.dumps([{'id': case.id, 'input_data': '2', 'expected_output': '2', 'is_sample': True}]),
        })
        self.assertRedirects(response, reverse('staff:problem_detail', args=[problem.id]))
        problem.refresh_from_db()
        self.assertEqual(problem.title, 'Managed problem')
        self.assertEqual(problem.test_cases.get().input_data, '2')
        self.client.post(reverse('staff:test_case_create', args=[problem.id]), {'input_data': '3', 'expected_output': '3'})
        self.assertEqual(problem.test_cases.count(), 2)
        self.client.post(reverse('staff:test_case_delete', args=[problem.id, case.id]))
        self.assertEqual(problem.test_cases.count(), 1)
        self.client.get(reverse('staff:problem_delete', args=[problem.id]))
        self.assertTrue(Problem.objects.filter(id=problem.id).exists())
        self.client.post(reverse('staff:problem_delete', args=[problem.id]))
        self.assertFalse(Problem.objects.filter(id=problem.id).exists())

    def test_staff_can_add_multiline_expected_output(self):
        self.client.force_login(self.staff)
        problem = Problem.objects.create(title='Multiline case problem', description='d', difficulty='easy')
        expected = 'first line\nsecond line\nthird line'

        response = self.client.post(reverse('staff:test_case_create', args=[problem.id]), {
            'input_data': 'value',
            'expected_output': expected,
        })

        self.assertRedirects(response, reverse('staff:problem_detail', args=[problem.id]))
        self.assertEqual(problem.test_cases.get().expected_output, expected)
